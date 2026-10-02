from docx import Document

doc = Document(r'd:\Neel College\Projects\Data analytics project\DATA_nova.docx')

with open(r'd:\Neel College\Projects\Data analytics project\DATA_nova_extracted.txt', 'w', encoding='utf-8') as f:
    for para in doc.paragraphs:
        f.write(para.text + '\n')
    
    # Also extract tables
    for i, table in enumerate(doc.tables):
        f.write(f'\n--- TABLE {i+1} ---\n')
        for row in table.rows:
            row_data = [cell.text for cell in row.cells]
            f.write(' | '.join(row_data) + '\n')

print("Extraction complete!")
print(f"Paragraphs: {len(doc.paragraphs)}")
print(f"Tables: {len(doc.tables)}")
